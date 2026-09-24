"""Выгрузка справочников из шаблона доменных данных в CSV.

Шаблон (xlsx) в репозиторий не кладётся: он рабочий, в нём есть листы,
не предназначенные для публикации. В репозиторий попадают только листы,
перечисленные в SHEETS, в виде CSV — это и есть справочники, из которых
загрузчики строят данные мока учётной системы и платформы.

Выгрузка детерминирована: при неизменном шаблоне повторный запуск даёт
побайтно те же файлы. Колонки сохраняются под исходными именами, текст —
дословно: предписанные действия правил (ADR-0004, раздел 5) не должны
искажаться ни при каком преобразовании.

    python data/generators/export_reference.py [путь/к/шаблону.xlsx]
"""

from __future__ import annotations

import csv
import datetime as dt
import hashlib
import json
import re
import sys
from pathlib import Path

import openpyxl

GENERATOR_VERSION = "1"

REPO_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_SOURCE = REPO_ROOT / "Шаблон_данных_проекта_s_3_региона.xlsx"
OUT_DIR = REPO_ROOT / "data" / "reference"

# Лист шаблона -> имя файла. Исходные клиентские записи и спецификация
# конкретного поставщика в публичный набор не выгружаются.
SHEETS = {
    "01_Модели_оборудования": "01_models.csv",
    "02_Коды_ошибок": "02_error_codes.csv",
    "03_Компетенции": "03_skills.csv",
    "04_Правила_решения": "04_decision_rules.csv",
    "05_Правила_гарантии": "05_warranty_rules.csv",
    "06_Правила_согласования": "06_approval_rules.csv",
    "08_Пользователи_и_роли": "08_users.csv",
    "09_Сервисные_организации": "09_service_organizations.csv",
    "10_Документы": "10_documents.csv",
    "11_Клиенты": "11_partners.csv",
    "12_Оборудование": "12_equipment.csv",
    "13_Договоры": "13_contracts.csv",
    "14_Инженеры": "14_technicians.csv",
    "15_Запчасти": "15_parts.csv",
    "16_История_обращений": "16_service_history.csv",
    "18_Сервисные_центры": "18_service_centers.csv",
    "19_Доступ_к_территориям": "19_territory_access.csv",
}

# Строка заголовка — первая, где первая ячейка является машинным
# идентификатором колонки (model_id, rule_id, ...). Строки выше — пояснения
# к заполнению шаблона.
HEADER_RE = re.compile(r"^[a-z_]+$")


def cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, dt.datetime):
        return value.date().isoformat() if value.time() == dt.time() else value.isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def export_sheet(ws) -> tuple[list[str], list[list[str]]]:
    rows = [[cell_text(c) for c in r] for r in ws.iter_rows(values_only=True)]
    rows = [r for r in rows if any(r)]
    header_at = next(i for i, r in enumerate(rows) if HEADER_RE.match(r[0]))
    header = rows[header_at]
    width = max(i for i, h in enumerate(header) if h) + 1
    header = header[:width]
    data = [r[:width] for r in rows[header_at + 1:] if any(r[:width])]
    return header, data


def main() -> None:
    source = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_SOURCE
    if not source.exists():
        sys.exit(f"Шаблон не найден: {source}")

    wb = openpyxl.load_workbook(source, read_only=True, data_only=True)
    missing = [s for s in SHEETS if s not in wb.sheetnames]
    if missing:
        sys.exit(f"В шаблоне нет листов: {', '.join(missing)}")

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    manifest = {
        "generator": "data/generators/export_reference.py",
        "generator_version": GENERATOR_VERSION,
        "source_file": source.name,
        "source_sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
        "files": {},
    }
    for sheet, filename in SHEETS.items():
        header, data = export_sheet(wb[sheet])
        with (OUT_DIR / filename).open("w", encoding="utf-8", newline="") as f:
            writer = csv.writer(f, lineterminator="\n")
            writer.writerow(header)
            writer.writerows(data)
        manifest["files"][filename] = {"sheet": sheet, "rows": len(data), "columns": header}
        print(f"{filename:32s} {len(data):4d} строк  <- {sheet}")

    (OUT_DIR / "manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )


if __name__ == "__main__":
    main()
